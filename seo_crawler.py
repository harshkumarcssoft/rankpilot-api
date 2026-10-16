import argparse, csv, fnmatch, hashlib, ipaddress, json, math, re, socket, time, xml.etree.ElementTree as ET
from collections import deque, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse, urljoin, urldefrag, urlencode, parse_qsl, urlunparse, quote
from urllib.robotparser import RobotFileParser

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADERS = ['Address', 'Content Type', 'Status Code', 'Status', 'Indexability', 'Indexability Status', 'Title 1', 'Title 1 Length', 'Title 1 Pixel Width', 'Meta Description 1', 'Meta Description 1 Length', 'Meta Description 1 Pixel Width', 'Meta Keywords 1', 'Meta Keywords 1 Length', 'H1-1', 'H1-1 Length', 'H1-2', 'H1-2 Length', 'H2-1', 'H2-1 Length', 'H2-2', 'H2-2 Length', 'Meta Robots 1', 'X-Robots-Tag 1', 'Meta Refresh 1', 'Canonical Link Element 1', 'rel="next" 1', 'rel="prev" 1', 'HTTP rel="next" 1', 'HTTP rel="prev" 1', 'amphtml Link Element', 'Size (bytes)', 'Transferred (bytes)', 'Total Transferred (bytes)', 'CO2 (mg)', 'Carbon Rating', 'Word Count', 'Sentence Count', 'Average Words Per Sentence', 'Flesch Reading Ease Score', 'Readability', 'Text Ratio', 'Crawl Depth', 'Folder Depth', 'Link Score', 'Inlinks', 'Unique Inlinks', 'Unique JS Inlinks', '% of Total', 'Outlinks', 'Unique Outlinks', 'Unique JS Outlinks', 'External Outlinks', 'Unique External Outlinks', 'Unique External JS Outlinks', 'Closest Similarity Match', 'No. Near Duplicates', 'Spelling Errors', 'Grammar Errors', 'Hash', 'Response Time', 'Last Modified', 'Redirect URL', 'Redirect Type', 'Cookies', 'Language', 'HTTP Version', 'Extractor 1 1', 'Extractor 2 1', 'Extractor 2 2', 'Extractor 3 1', 'Extractor 3 2', 'Mobile Alternate Link', 'URL Encoded Address', 'Crawl Timestamp']
UA = "SiteAuditCrawler/1.0 (+respectful technical SEO audit)"
TIMEOUT = (8, 18)
MAX_BYTES = 8_000_000
TRACKING = {"utm_source","utm_medium","utm_campaign","utm_term","utm_content","gclid","fbclid","msclkid"}
NON_HTML = (".jpg",".jpeg",".png",".gif",".webp",".svg",".ico",".pdf",".zip",".gz",".css",".js",".mp4",".mp3",".woff",".woff2",".ttf",".eot")

def utcnow():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")

def public_host(host):
    if not host or host.lower() in {"localhost","localhost.localdomain"} or host.endswith(".local"):
        return False
    try:
        ips=[ipaddress.ip_address(host)]
    except ValueError:
        try:
            ips=[ipaddress.ip_address(x[4][0]) for x in socket.getaddrinfo(host,None,type=socket.SOCK_STREAM)]
        except Exception:
            return False
    return all(ip.is_global for ip in ips)

def host_family(host):
    h=(host or "").lower().rstrip(".")
    return h[4:] if h.startswith("www.") else h

def normalize(raw):
    try:
        raw=urldefrag(str(raw).strip())[0]
        p=urlparse(raw)
        if p.scheme.lower() not in ("http","https") or not p.hostname or not public_host(p.hostname):
            return None
        scheme=p.scheme.lower(); host=p.hostname.lower().rstrip("."); port=p.port
        if (scheme=="http" and port==80) or (scheme=="https" and port==443): port=None
        netloc=host if port is None else f"{host}:{port}"
        path=p.path or "/"
        query=urlencode([(k,v) for k,v in parse_qsl(p.query,keep_blank_values=True) if k.lower() not in TRACKING],doseq=True)
        return urlunparse((scheme,netloc,path,"",query,""))
    except Exception:
        return None

def same_site(url, family):
    try: return host_family(urlparse(url).hostname)==family
    except Exception: return False

def txt(node):
    return node.get_text(" ",strip=True) if node else ""

def meta_values(soup, name):
    """Return every matching meta value, preserving document order and empty values."""
    values=[]
    for tag in soup.find_all("meta", attrs={"name": re.compile(r"^\s*"+re.escape(name)+r"\s*$", re.I)}):
        value=tag.get("content")
        values.append(str(value).strip() if value is not None else "")
    return values

def meta(soup, name):
    """Use the first non-empty standard meta[name] value, not an empty duplicate."""
    return next((value for value in meta_values(soup, name) if value), "")

def looks_like_html(body):
    """Conservative HTML sniffing for servers with missing/wrong Content-Type headers."""
    sample=(body or b"")[:65536].lstrip().lower()
    if sample.startswith((b"<!doctype html", b"<html", b"<?xml")) and (b"<html" in sample or b"xhtml" in sample):
        return True
    return bool(re.search(rb"<(?:html|head|body|title|meta|!doctype\s+html)(?:\s|>|/)", sample))

def is_html_response(content_type, body):
    c=(content_type or "").lower()
    return "text/html" in c or "application/xhtml+xml" in c or looks_like_html(body)

def rel_link(soup, rel, final):
    for t in soup.find_all("link",href=True):
        rs=t.get("rel",[]); rs=[rs] if isinstance(rs,str) else rs
        if rel.lower() in [str(x).lower() for x in rs]:
            return urljoin(final,t["href"])
    return ""

def pixel_width(text):
    # Approximation based on common sans-serif glyph widths; not a browser-rendered measurement.
    widths={"i":4,"l":4,"I":4,"t":5,"f":5,"r":5,"m":11,"w":10,"M":11,"W":12," ":4}
    return round(sum(widths.get(c,7) for c in (text or "")))

def flesch(text, words, sentences):
    if words == 0 or sentences == 0: return ""
    def syllables(word):
        w=re.sub(r"[^a-z]","",word.lower())
        if not w: return 0
        groups=re.findall(r"[aeiouy]+",w)
        count=len(groups)
        if w.endswith("e") and count>1 and not w.endswith(("le","ye")): count-=1
        return max(count,1)
    total=sum(syllables(w) for w in re.findall(r"\b[\w'-]+\b",text.lower()))
    score=206.835-1.015*(words/sentences)-84.6*(total/words)
    return round(score,2)

def readability(score):
    if score=="": return ""
    if score>=90: return "Very Easy"
    if score>=80: return "Easy"
    if score>=70: return "Fairly Easy"
    if score>=60: return "Standard"
    if score>=50: return "Fairly Difficult"
    if score>=30: return "Difficult"
    return "Very Difficult"

def fetch_xml_or_feed_urls(session, seed_url, family, limit=100000, source_diagnostics=None, source_kind="sitemap"):
    """Read sitemap indexes and feeds; optionally record every source outcome."""
    queue=deque([(seed_url, "sitemap")]); seen=set(); found=[]
    while queue and len(found)<limit:
        sm, inherited_source=queue.popleft(); n=normalize(sm)
        if not n or n in seen: continue
        seen.add(n)
        try:
            r=session.get(n,timeout=TIMEOUT,allow_redirects=True,headers={"User-Agent":UA},stream=True)
            status_code=r.status_code; content_type=r.headers.get("Content-Type", "")
            if status_code>=400:
                if source_diagnostics is not None:
                    source_diagnostics.append({"Source type":source_kind,"Source URL":n,"HTTP status":status_code,"URLs discovered":0,"Source status":"blocked_or_failed","Notes":f"HTTP {status_code}; source content was not available"})
                r.close(); continue
            content=r.raw.read(MAX_BYTES,decode_content=True); r.close()
            before=len(found); before_queue=len(queue); root=ET.fromstring(content)
            root_name=root.tag.rsplit('}',1)[-1].lower()
            if root_name == 'sitemapindex':
                for el in root.iter():
                    if el.tag.rsplit('}',1)[-1].lower()=='loc' and el.text:
                        queue.append((el.text.strip(),'sitemap_index'))
            elif root_name in ('urlset','rss','feed','rdf'):
                if root_name=='urlset':
                    for el in root.iter():
                        if el.tag.rsplit('}',1)[-1].lower()=='loc' and el.text:
                            u=normalize(urljoin(n,el.text.strip()))
                            if u and same_site(u,family): found.append((u,'sitemap'))
                else:
                    for el in root.iter():
                        tag=el.tag.rsplit('}',1)[-1].lower()
                        val=el.attrib.get('href') if tag=='link' else (el.text.strip() if tag in ('link','guid','id') and el.text else '')
                        if val:
                            if tag=='link' and el.attrib.get('rel') not in (None,'alternate'): continue
                            u=normalize(urljoin(n,val))
                            if u and same_site(u,family): found.append((u,'feed'))
            if source_diagnostics is not None:
                discovered_delta=(len(found)-before)+(len(queue)-before_queue)
                source_diagnostics.append({"Source type":source_kind,"Source URL":n,"HTTP status":status_code,"URLs discovered":discovered_delta,"Source status":"checked","Content type":content_type,"Notes":"Parsed successfully"})
            if len(found)>=limit: break
        except (requests.RequestException, ET.ParseError, ValueError) as exc:
            if source_diagnostics is not None:
                source_diagnostics.append({"Source type":source_kind,"Source URL":n,"HTTP status":"","URLs discovered":0,"Source status":"failed","Notes":f"{type(exc).__name__}: {exc}"})
            continue
    return found[:limit]

def sitemap_urls(session, origin, family, limit, robots_text='', source_diagnostics=None):
    seeds=[]
    for line in (robots_text or '').splitlines():
        m=re.match(r'\s*sitemap\s*:\s*(\S+)',line,re.I)
        if m: seeds.append(m.group(1))
    seeds.extend([urljoin(origin,'/sitemap.xml'),urljoin(origin,'/sitemap_index.xml')])
    out=[]; seen=set()
    for seed in dict.fromkeys(seeds):
        for u,source in fetch_xml_or_feed_urls(session,seed,family,max(0,limit-len(out)),source_diagnostics=source_diagnostics,source_kind='sitemap'):
            if u not in seen: seen.add(u); out.append(u)
        if len(out)>=limit: break
    return out

def load_url_inventory(path, start, family, source_name):
    """Load URLs from a plain text/CSV file or an earlier Rankpilot workbook."""
    out=[]; path=Path(path)
    if not path.exists(): return out
    try:
        if path.suffix.lower() in ('.xlsx','.xlsm'):
            from openpyxl import load_workbook
            wb=load_workbook(path,read_only=True,data_only=True)
            for ws in wb.worksheets:
                headers=[str(c or '').strip().lower() for c in next(ws.iter_rows(values_only=True),())]
                col=next((i for i,h in enumerate(headers) if h in ('address','url','requested url','final url')),None)
                if col is None: continue
                for values in ws.iter_rows(min_row=2,values_only=True):
                    if col<len(values) and values[col]:
                        u=normalize(urljoin(start,str(values[col]).strip()))
                        if u and same_site(u,family): out.append((u,source_name))
                if out: break
            wb.close()
        else:
            with path.open('r',encoding='utf-8-sig',errors='replace',newline='') as f:
                if path.suffix.lower()=='.csv':
                    for row in csv.reader(f):
                        for value in row:
                            if value and re.match(r'^https?://',value.strip(),re.I):
                                u=normalize(value.strip())
                                if u and same_site(u,family): out.append((u,source_name))
                else:
                    for line in f:
                        value=line.strip().split()[0] if line.strip() else ''
                        u=normalize(urljoin(start,value)) if value else None
                        if u and same_site(u,family): out.append((u,source_name))
    except Exception:
        return out
    return out

def extract_extra_page_data(url, final_url, response, body, soup, links, rendered_body=None):
    """Extended page-level evidence. Blank means unavailable/not observable, not a positive pass."""
    raw_soup=BeautifulSoup(body,'lxml') if is_html_response(response.headers.get('Content-Type',''),body) else None
    text=txt(soup.body or soup) if soup else ''
    visible_words=re.findall(r"\b[\w'-]+\b",text)
    images=[]; links_data=[]; resources=[]; schema=[]; forms=[]
    if soup:
        for im in soup.find_all('img'):
            src=im.get('src') or im.get('data-src') or ''
            images.append({'src':urljoin(final_url,src) if src else '', 'alt':im.get('alt'), 'alt_present':im.has_attr('alt'), 'width':im.get('width',''), 'height':im.get('height',''), 'loading':im.get('loading',''), 'srcset':im.get('srcset',''), 'title':im.get('title','')})
        for a in soup.find_all('a',href=True):
            target=urljoin(final_url,a['href']); label=txt(a) or a.get('aria-label','') or a.get('title','')
            links_data.append({'url':target,'text':label,'rel':' '.join(a.get('rel',[])) if isinstance(a.get('rel',[]),list) else str(a.get('rel','')),'target':a.get('target','')})
        for el in soup.find_all(['script','link','iframe','source'],src=True)+soup.find_all('link',href=True):
            attr='src' if el.has_attr('src') else 'href'; raw=el.get(attr,'')
            if raw: resources.append({'tag':el.name,'url':urljoin(final_url,raw),'type':el.get('type',''),'async':el.has_attr('async'),'defer':el.has_attr('defer')})
        for tag in soup.find_all('script',attrs={'type':re.compile(r'application/ld\+json',re.I)}):
            raw=tag.string or tag.get_text()
            try: schema.append(json.loads(raw))
            except Exception: schema.append({'parse_error':True,'raw_excerpt':raw[:500]})
        for form in soup.find_all('form'):
            forms.append({'action':urljoin(final_url,form.get('action','')),'method':form.get('method','get').lower(),'fields':len(form.find_all(['input','select','textarea','button'])),'has_labelled_fields':all((x.get('aria-label') or x.get('id') and soup.find('label',attrs={'for':x.get('id')})) for x in form.find_all(['input','select','textarea']) if x.get('type') not in ('hidden','submit','button'))})
    meta_all=[]
    if soup:
        for tag in soup.find_all('meta'):
            meta_all.append({'name':tag.get('name',''),'property':tag.get('property',''),'http_equiv':tag.get('http-equiv',''),'content':tag.get('content','')})
    headings=[]
    if soup:
        headings=[{'level':int(tag.name[1]),'text':txt(tag)} for tag in soup.find_all(re.compile(r'^h[1-6]$'))]
    canonicals=[]
    if soup:
        for tag in soup.find_all('link',href=True):
            rel=tag.get('rel',[]); rel=[rel] if isinstance(rel,str) else rel
            if 'canonical' in [str(x).lower() for x in rel]: canonicals.append(urljoin(final_url,tag['href']))
    robots_meta=[]
    if soup:
        for tag in soup.find_all('meta',attrs={'name':re.compile(r'^(robots|googlebot|bingbot)$',re.I)}): robots_meta.append({'name':tag.get('name',''),'content':tag.get('content','')})
    hreflang_links=[]
    if soup:
        for tag in soup.find_all('link',href=True):
            rel=tag.get('rel',[]); rel=[rel] if isinstance(rel,str) else rel
            if 'alternate' in [str(x).lower() for x in rel] and tag.get('hreflang'):
                hreflang_links.append({'hreflang':tag.get('hreflang'),'url':urljoin(final_url,tag['href'])})
    accessibility_signals={
        'images_total':len(images),
        'images_missing_alt':sum(1 for im in images if not im.get('alt_present') or not str(im.get('alt') or '').strip()),
        'forms_total':len(forms),
        'forms_with_unlabelled_fields':sum(1 for form in forms if not form.get('has_labelled_fields')),
        'html_lang_missing':not bool(soup.html.get('lang','').strip()) if soup and soup.html else None,
        'viewport_meta_present':bool(soup and soup.find('meta',attrs={'name':re.compile(r'^viewport$',re.I)})),
        'h1_count':sum(1 for h in headings if h['level']==1),
        'heading_order_observation':'See Headings JSON for full heading sequence; automated accessibility conformance is not claimed'
    }
    local_seo_signals={
        'business_schema_present':any(any(k in str(obj.get('@type','')).lower() for k in ('localbusiness','organization','medicalbusiness','dentist')) for obj in schema if isinstance(obj,dict)),
        'address_text_present':bool(soup and re.search(r'\b(address|street|postal|zip code)\b',text,re.I)),
        'phone_text_present':bool(re.search(r'(?:\+?\d[\d().\s-]{7,}\d)',text)),
        'map_link_present':bool(soup and any('maps.google.' in str(a.get('href','')).lower() or 'maps.app.goo.gl' in str(a.get('href','')).lower() for a in soup.find_all('a',href=True))),
        'note':'Presence signals only; they do not verify business identity or NAP consistency'
    }
    return {
      'Address':url,'Final URL':final_url,'HTTP status':response.status_code,'Content type':response.headers.get('Content-Type',''),
      'Response headers JSON':json.dumps(dict(response.headers),ensure_ascii=False),'Fetch timestamp UTC':utcnow(),
      'Raw HTML bytes':len(body),'Body truncated':len(body)>=MAX_BYTES,'HTML detected':bool(soup),
      'Rendered DOM checked':bool(rendered_body),'Raw vs rendered': 'rendered_and_raw' if rendered_body else 'raw_only',
      'Title tags JSON':json.dumps([txt(t) for t in soup.find_all('title')] if soup else [],ensure_ascii=False),
      'Meta tags JSON':json.dumps(meta_all,ensure_ascii=False),'Headings JSON':json.dumps(headings,ensure_ascii=False),
      'H1 count':sum(1 for h in headings if h['level']==1),'H2 count':sum(1 for h in headings if h['level']==2),
      'Canonical values JSON':json.dumps(canonicals,ensure_ascii=False),'Robots directives JSON':json.dumps(robots_meta,ensure_ascii=False),'Hreflang links JSON':json.dumps(hreflang_links,ensure_ascii=False),
      'X-Robots-Tag':response.headers.get('X-Robots-Tag',''),'Visible text':text,'Word count':len(visible_words),
      'Character count':len(text),'Paragraph count':len(soup.find_all('p')) if soup else 0,
      'Image data JSON':json.dumps(images,ensure_ascii=False),'Links data JSON':json.dumps(links_data,ensure_ascii=False),
      'Structured data JSON':json.dumps(schema,ensure_ascii=False),'Resources JSON':json.dumps(resources,ensure_ascii=False),
      'Forms JSON':json.dumps(forms,ensure_ascii=False),'HTML lang':soup.html.get('lang','') if soup and soup.html else '',
      'Viewport':next((m.get('content','') for m in soup.find_all('meta',attrs={'name':re.compile(r'^viewport$',re.I)})), '') if soup else '',
      'Open Graph JSON':json.dumps([m for m in meta_all if str(m.get('property','')).lower().startswith('og:')],ensure_ascii=False),
      'Twitter card JSON':json.dumps([m for m in meta_all if str(m.get('name','')).lower().startswith('twitter:')],ensure_ascii=False),
      'Internal links found':sum(1 for l in links_data if same_site(l['url'],host_family(urlparse(url).hostname))),
      'External links found':sum(1 for l in links_data if not same_site(l['url'],host_family(urlparse(url).hostname))),
      'Response security headers JSON':json.dumps({h:response.headers.get(h,'') for h in ['Strict-Transport-Security','Content-Security-Policy','X-Content-Type-Options','Referrer-Policy','Permissions-Policy','X-Frame-Options']},ensure_ascii=False),
      'Accessibility signals JSON':json.dumps(accessibility_signals,ensure_ascii=False),'Local SEO presence signals JSON':json.dumps(local_seo_signals,ensure_ascii=False),
      'Performance measurement note':'HTTP response time and transferred bytes are measured; Core Web Vitals/Lighthouse require a separate authorised measurement integration',
      'Extraction status':'checked' if soup and response.status_code<400 and len(body)<MAX_BYTES else ('incomplete' if len(body)>=MAX_BYTES else 'not_applicable')
    }

def is_challenge_response(status_code, body, headers=None):
    """Detect common access challenges; never treat challenge HTML as page SEO content."""
    if int(status_code or 0) not in (401, 403, 429, 503):
        return False
    sample=(body or b"")[:250000].decode("utf-8", errors="ignore").lower()
    header_text=" ".join(f"{k}:{v}" for k,v in (headers or {}).items()).lower()
    markers=("just a moment", "cf-chl-", "challenge-platform", "checking your browser",
             "verify you are human", "attention required! | cloudflare", "cf-ray")
    return any(marker in sample or marker in header_text for marker in markers)


def path_matches_rules(url, patterns):
    """Match include/exclude globs against URL path or full URL."""
    parsed=urlparse(url); path=parsed.path or "/"
    for pattern in patterns or []:
        pat=str(pattern).strip()
        if not pat or pat.startswith("#"):
            continue
        if fnmatch.fnmatch(path,pat) or fnmatch.fnmatch(url,pat):
            return True
    return False


def extract_page(url, response, final_url, chain, body, depth, all_rows, inlinks, reported_status=None, rendered_body=None):
    ctype=response.headers.get("Content-Type","")
    content_length=response.headers.get("Content-Length","")
    try:
        declared_size=int(content_length)
    except (TypeError,ValueError):
        declared_size=len(body)
    body_truncated=declared_size > len(body) or len(body) >= MAX_BYTES
    body_hash=hashlib.sha256(body).hexdigest() if not body_truncated else ""
    html_detected=is_html_response(ctype, body)
    raw_soup=BeautifulSoup(body,"lxml") if html_detected else None
    rendered_soup=BeautifulSoup(rendered_body,"lxml") if rendered_body else None
    # Use rendered DOM when explicitly requested and available; keep raw extraction evidence separately.
    soup=rendered_soup or raw_soup
    raw_title_values=[txt(t) for t in raw_soup.find_all("title")] if raw_soup else []
    raw_desc_values=meta_values(raw_soup,"description") if raw_soup else []
    raw_h1_values=[txt(t) for t in raw_soup.find_all("h1")] if raw_soup else []
    selected_title_values=[txt(t) for t in soup.find_all("title")] if soup else []
    selected_desc_values=meta_values(soup,"description") if soup else []
    selected_h1_values=[txt(t) for t in soup.find_all("h1")] if soup else []
    extraction_status=("not_applicable" if not html_detected or not (200 <= (reported_status if reported_status is not None else response.status_code) < 300)
                       else ("incomplete" if body_truncated and not rendered_body else "checked"))
    diagnostics={"Address":url,"status_code":reported_status if reported_status is not None else response.status_code,
        "content_type":ctype,"html_detected":bool(html_detected),"extraction_status":extraction_status,
        "body_truncated":bool(body_truncated),"raw_bytes_read":len(body),"declared_bytes":declared_size,
        "rendered_checked":bool(rendered_body),"title_tags_raw":len(raw_title_values),
        "meta_description_tags_raw":len(raw_desc_values),"h1_tags_raw":len(raw_h1_values),
        "title_raw":next((x for x in raw_title_values if x),""),
        "meta_description_raw":next((x for x in raw_desc_values if x),""),
        "h1_raw":next((x for x in raw_h1_values if x),""),
        "title_selected":next((x for x in selected_title_values if x),""),
        "meta_description_selected":next((x for x in selected_desc_values if x),""),
        "h1_selected":next((x for x in selected_h1_values if x),""),
        "open_graph_description_present":bool(raw_soup and raw_soup.find("meta",attrs={"property":re.compile(r"^og:description$",re.I)})),
        "note":"Rendered DOM used for primary fields" if rendered_body else "Raw response HTML only; JavaScript rendering was not performed"}
    row={h:"" for h in HEADERS}
    status_code = reported_status if reported_status is not None else response.status_code
    status_label = "Success" if 200 <= status_code < 300 else ("Redirect" if 300 <= status_code < 400 else ("Client Error" if 400 <= status_code < 500 else ("Server Error" if status_code >= 500 else "Unknown")))
    row.update({"Address":url,"Content Type":ctype,"Status Code":status_code,
                "Status": status_label,
                "Size (bytes)":declared_size,"Transferred (bytes)":len(body),
                "Total Transferred (bytes)":len(body),"Response Time":round(response.elapsed.total_seconds(),3),
                "Last Modified":response.headers.get("Last-Modified",""),"Redirect URL":final_url if final_url!=url else "",
                "Redirect Type": " -> ".join(chain) if len(chain)>1 else "",
                "Cookies":"; ".join(response.headers.get("Set-Cookie","").splitlines()),
                "HTTP Version":getattr(response.raw,"version",""),
                "URL Encoded Address":quote(url,safe=":/?&=%#[]@!$'()*+,;"),
                "Crawl Timestamp":utcnow(),"Crawl Depth":depth,
                "Folder Depth":len([x for x in urlparse(url).path.split("/") if x]),
                "Inlinks":len(inlinks.get(url,set())),"Unique Inlinks":len(inlinks.get(url,set())),
                "Unique JS Inlinks":"","% of Total":"","Link Score":"",
                "Outlinks":0,"Unique Outlinks":0,"Unique JS Outlinks":"",
                "External Outlinks":0,"Unique External Outlinks":0,"Unique External JS Outlinks":"",
                "Closest Similarity Match":"","No. Near Duplicates":0,
                "Spelling Errors":"","Grammar Errors":"","Hash":body_hash})
    challenge=is_challenge_response(status_code, body, response.headers)
    if status_code >= 400 or challenge:
        row["Indexability"]="Blocked" if challenge or status_code in (401,403,429) else "Non-Indexable"
        row["Indexability Status"]=("Access challenge detected; page content not assessed" if challenge else f"HTTP {status_code}; page content not assessed")
        diagnostics.update({"extraction_status":"blocked_challenge" if challenge else "not_applicable",
            "challenge_detected":bool(challenge),
            "note":"Access challenge detected; no SEO metadata or links were extracted from this response" if challenge else f"HTTP {status_code}; SEO metadata and links were not assessed from an error response"})
        return row, [], diagnostics
    links=[]
    if not soup:
        row["Indexability"]="Non-HTML"
        row["Indexability Status"]="Not HTML or HTML could not be detected from response"
        diagnostics["extraction_status"]="not_applicable"
        return row,links,diagnostics
    title=next((txt(t) for t in soup.find_all("title") if txt(t)), "")
    desc=meta(soup,"description"); keywords=meta(soup,"keywords")
    h1=[txt(x) for x in soup.find_all("h1") if txt(x)]; h2=[txt(x) for x in soup.find_all("h2") if txt(x)]
    robots=meta(soup,"robots")
    xrobots=response.headers.get("X-Robots-Tag","")
    noindex=bool(re.search(r"(?<![\w-])noindex(?![\w-])",(robots+" "+xrobots).lower()))
    if noindex:
        indexability="Non-Indexable"; indexability_status="Blocked by noindex"
    elif 300 <= status_code < 400:
        indexability="Non-Indexable"; indexability_status="Redirect response"
    elif status_code >= 400:
        indexability="Non-Indexable"; indexability_status=f"HTTP {status_code} response"
    elif 200 <= status_code < 300:
        indexability="Indexable (directive check only)"; indexability_status="No noindex detected"
    else:
        indexability="Unknown"; indexability_status="HTTP status not confirmed"
    canonical=rel_link(soup,"canonical",final_url)
    next_link=rel_link(soup,"next",final_url); prev_link=rel_link(soup,"prev",final_url)
    refresh=next((t.get("content","") for t in soup.find_all("meta",attrs={"http-equiv":re.compile("^refresh$",re.I)})),"")
    text=txt(soup.body or soup)
    words=re.findall(r"\b[\w'-]+\b",text)
    sentences=[s for s in re.split(r"[.!?]+",text) if s.strip()]
    avg=round(len(words)/len(sentences),2) if sentences else ""
    score=flesch(text,len(words),len(sentences))
    text_ratio=round(len(text)/max(len(body),1),4)
    lang=(soup.html.get("lang","") if soup.html else "")
    images=soup.find_all("img")
    link_entries=[]
    for a in soup.find_all("a",href=True):
        target=normalize(urljoin(final_url,a["href"]))
        if target:
            internal=host_family(urlparse(target).hostname)==host_family(urlparse(url).hostname)
            link_entries.append((target,internal))
            links.append(target)
    internal=[u for u,i in link_entries if i]; external=[u for u,i in link_entries if not i]
    jsonld=[]
    for tag in soup.find_all("script",attrs={"type":re.compile("application/ld\\+json",re.I)}):
        try:
            obj=json.loads(tag.string or tag.get_text())
            for item in (obj if isinstance(obj,list) else [obj]):
                if isinstance(item,dict) and item.get("@type"):
                    jsonld.extend(item["@type"] if isinstance(item["@type"],list) else [item["@type"]])
        except Exception: pass
    alternates=[]
    for t in soup.find_all("link",href=True):
        rs=t.get("rel",[]); rs=[rs] if isinstance(rs,str) else rs
        if "alternate" in rs and t.get("hreflang"): alternates.append(t.get("hreflang")+"="+urljoin(final_url,t["href"]))
    h1s=h1+[""]*2; h2s=h2+[""]*2
    row.update({
        "Indexability":indexability,
        "Indexability Status":indexability_status,
        "Title 1":title,"Title 1 Length":len(title),"Title 1 Pixel Width":pixel_width(title),
        "Meta Description 1":desc,"Meta Description 1 Length":len(desc),
        "Meta Description 1 Pixel Width":pixel_width(desc),"Meta Keywords 1":keywords,
        "Meta Keywords 1 Length":len(keywords),
        "H1-1":h1s[0],"H1-1 Length":len(h1s[0]),"H1-2":h1s[1],"H1-2 Length":len(h1s[1]),
        "H2-1":h2s[0],"H2-1 Length":len(h2s[0]),"H2-2":h2s[1],"H2-2 Length":len(h2s[1]),
        "Meta Robots 1":robots,"X-Robots-Tag 1":xrobots,"Meta Refresh 1":refresh,
        "Canonical Link Element 1":canonical,'rel="next" 1':next_link,'rel="prev" 1':prev_link,
        "amphtml Link Element":rel_link(soup,"amphtml",final_url),
        "Word Count":len(words),"Sentence Count":len(sentences),
        "Average Words Per Sentence":avg,"Flesch Reading Ease Score":score,
        "Readability":readability(score) if score!="" else "","Text Ratio":text_ratio,
        "Outlinks":len(link_entries),"Unique Outlinks":len(set(links)),
        "External Outlinks":len(external),"Unique External Outlinks":len(set(external)),
        "Language":lang,"Extractor 1 1":", ".join(sorted(set(map(str,jsonld)))),
        "Mobile Alternate Link":next((urljoin(final_url,t["href"]) for t in soup.find_all("link",href=True)
            if "alternate" in (t.get("rel",[]) if isinstance(t.get("rel",[]),list) else [t.get("rel")]) and t.get("media")),""),
    })
    return row,links,diagnostics

def write_excel(rows, out, diagnostics=None, extra_rows=None, discovery_rows=None, discovery_diagnostics=None, crawl_summary=None):
    """Write every collected field; HEADERS is a baseline, not a column limit.

    New crawler fields are automatically appended to the main worksheet. Rich page
    extraction and diagnostic records also remain available in their dedicated sheets.
    """
    # Keep the familiar base fields first, then include every additional key found in
    # this crawl. This deliberately avoids a fixed column-count/schema ceiling.
    all_headers = list(HEADERS)
    for collection in (rows, list((extra_rows or {}).values()), list((diagnostics or {}).values())):
        for record in collection:
            for key in record.keys():
                key = str(key)
                if key not in all_headers:
                    all_headers.append(key)
    # Add the extended page-data fields to the main report with a prefix so fields
    # such as "Title" or "Language" cannot overwrite legacy fields with similar names.
    page_data_headers = []
    for record in list((extra_rows or {}).values()):
        for key in record.keys():
            column = "Page Data: " + str(key)
            if column not in all_headers:
                all_headers.append(column)
            if column not in page_data_headers:
                page_data_headers.append(column)
    diagnostic_headers = []
    for record in list((diagnostics or {}).values()):
        for key in record.keys():
            column = "Diagnostic: " + str(key)
            if column not in all_headers:
                all_headers.append(column)
            if column not in diagnostic_headers:
                diagnostic_headers.append(column)
    extras_by_url = extra_rows or {}
    diagnostics_by_url = diagnostics or {}
    wb=Workbook(); ws=wb.active; ws.title="1 - All"
    ws.append(all_headers)
    for row in rows:
        address = row.get("Address", "")
        combined = dict(row)
        for key, value in extras_by_url.get(address, {}).items():
            combined["Page Data: " + str(key)] = value
        for key, value in diagnostics_by_url.get(address, {}).items():
            combined["Diagnostic: " + str(key)] = value
        ws.append([combined.get(h, "") for h in all_headers])
    fill=PatternFill("solid",fgColor="17365D")
    for c in ws[1]:
        c.font=Font(color="FFFFFF",bold=True); c.fill=fill
        c.alignment=Alignment(wrap_text=True,vertical="center")
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions; ws.row_dimensions[1].height=32
    for i,h in enumerate(HEADERS,1):
        width=min(max(len(h)+2,14),30)
        if h in ("Address","Canonical Link Element 1","Meta Description 1","Title 1","Redirect URL"): width=38
        ws.column_dimensions[get_column_letter(i)].width=width
    # Keep complete diagnostic evidence in its own sheet as well as in the dynamic main report.
    diag_rows=list((diagnostics or {}).values())
    if diag_rows:
        dws=wb.create_sheet("Audit Diagnostics")
        diag_headers=list(dict.fromkeys(k for item in diag_rows for k in item.keys()))
        dws.append(diag_headers)
        for item in diag_rows: dws.append([item.get(h, "") for h in diag_headers])
        for c in dws[1]:
            c.font=Font(color="FFFFFF",bold=True); c.fill=fill
            c.alignment=Alignment(wrap_text=True,vertical="center")
        dws.freeze_panes="A2"; dws.auto_filter.ref=dws.dimensions
        for i,h in enumerate(diag_headers,1): dws.column_dimensions[get_column_letter(i)].width=min(max(len(h)+2,14),42)
    for sheet_name, records in [('Page Data', list((extra_rows or {}).values())), ('URL Discovery', list(discovery_rows or [])), ('Discovery Diagnostics', list(discovery_diagnostics or [])), ('Crawl Summary', list(crawl_summary or []))]:
        if not records and sheet_name not in ('Discovery Diagnostics','Crawl Summary'): continue
        sh=wb.create_sheet(sheet_name)
        cols=list(dict.fromkeys(k for rec in records for k in rec.keys()))
        if not cols: cols=['Metric','Value'] if sheet_name=='Crawl Summary' else ['Source type','Source URL','HTTP status','URLs discovered','Source status','Notes']
        sh.append(cols)
        for rec in records: sh.append([rec.get(k,'') for k in cols])
        for c in sh[1]:
            c.font=Font(color='FFFFFF',bold=True); c.fill=fill; c.alignment=Alignment(wrap_text=True,vertical='center')
        sh.freeze_panes='A2'; sh.auto_filter.ref=sh.dimensions
        for i,h in enumerate(cols,1): sh.column_dimensions[get_column_letter(i)].width=min(max(len(str(h))+2,14),42)
    target=Path(out)
    target.parent.mkdir(parents=True,exist_ok=True)
    temporary=target.with_name(target.stem + '.writing.xlsx')
    wb.save(temporary)
    temporary.replace(target)

def write_progress(progress_path, *, status, crawled, max_pages, current_url="", message="", errors=0, queued=0):
    """Write a tiny JSON heartbeat for the dashboard; safe to update after each URL."""
    if not progress_path:
        return
    payload = {
        "status": status,
        "crawled": int(crawled),
        "max_pages": int(max_pages),
        "queued": int(queued),
        "current_url": str(current_url or ""),
        "message": str(message or ""),
        "errors": int(errors),
        "updated_at": utcnow(),
    }
    target = Path(progress_path)
    temporary = target.with_suffix(target.suffix + ".tmp")
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        temporary.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        temporary.replace(target)
    except OSError as exc:
        print(f"WARNING: Could not update progress file: {exc}", flush=True)


def render_page_html(url, timeout=25):
    """Optionally render a page in Chromium. Requires `pip install playwright` and `playwright install chromium`."""
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        return None, "Playwright is not installed. Install it with: python -m pip install playwright; python -m playwright install chromium"
    try:
        with sync_playwright() as p:
            browser=p.chromium.launch(headless=True)
            page=browser.new_page(user_agent=UA)
            page.goto(url, wait_until="domcontentloaded", timeout=timeout*1000)
            # Allow common client-side head/meta updates to settle without waiting indefinitely for analytics.
            page.wait_for_timeout(1200)
            html=page.content().encode("utf-8")
            browser.close()
            return html, ""
    except Exception as exc:
        return None, f"JavaScript rendering failed: {type(exc).__name__}: {exc}"

def run_crawl(start_url,max_pages,delay,output,use_sitemap=True,progress_path=None,render_js=False,url_files=None,previous_report=None,feed_urls=None,respect_robots=True,user_agent=UA,include_patterns=None,exclude_patterns=None):
    start=normalize(start_url)
    if not start: raise SystemExit("Please enter a full public URL starting with https:// or http://")
    family=host_family(urlparse(start).hostname)
    origin=f"{urlparse(start).scheme}://{urlparse(start).netloc}"
    ua=(str(user_agent or UA).strip() or UA)[:300]
    session=requests.Session(); session.headers.update({"User-Agent":ua,"Accept":"text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8"})
    robots=RobotFileParser(); source_diagnostics=[]; robots_text=""; robots_state="unknown"
    try:
        rr=session.get(origin+"/robots.txt",timeout=TIMEOUT,allow_redirects=True)
        if 200 <= rr.status_code < 300:
            robots_text=rr.text; robots.parse(robots_text.splitlines()); robots_state="checked"
            source_diagnostics.append({"Source type":"robots.txt","Source URL":rr.url,"HTTP status":rr.status_code,"URLs discovered":len(re.findall(r"(?im)^\s*sitemap\s*:",robots_text)),"Source status":"checked","Notes":"Robots directives parsed"})
        elif 400 <= rr.status_code < 500:
            robots.parse([]); robots_state="unavailable_4xx"
            source_diagnostics.append({"Source type":"robots.txt","Source URL":rr.url,"HTTP status":rr.status_code,"URLs discovered":0,"Source status":"unavailable","Notes":f"robots.txt returned HTTP {rr.status_code}; no directives could be read"})
        else:
            robots.parse(["User-agent: *", "Disallow: /"]); robots_state="unavailable_server_error"
            source_diagnostics.append({"Source type":"robots.txt","Source URL":rr.url,"HTTP status":rr.status_code,"URLs discovered":0,"Source status":"unavailable","Notes":f"robots.txt returned HTTP {rr.status_code}; crawling disallowed conservatively because rules could not be verified"})
        rr.close()
    except requests.RequestException as exc:
        robots.parse(["User-agent: *", "Disallow: /"]); robots_state="unreachable"
        source_diagnostics.append({"Source type":"robots.txt","Source URL":origin+"/robots.txt","HTTP status":"","URLs discovered":0,"Source status":"unavailable","Notes":f"Could not verify robots rules: {type(exc).__name__}: {exc}; crawling disallowed conservatively"})
    queue=deque(); queued=set(); visited=set(); depths={}; rows=[]; diagnostics={}; extra_rows={}; discovery_rows=[]; discovery_sources=defaultdict(set); incoming_total=defaultdict(int); incoming_sources=defaultdict(set); discovered_links={}; fetch_errors=0
    write_progress(progress_path, status="running", crawled=0, max_pages=max_pages, message="Starting crawl", queued=len(queue))
    include_rules=[str(x).strip() for x in (include_patterns or []) if str(x).strip()]
    exclude_rules=[str(x).strip() for x in (exclude_patterns or []) if str(x).strip()]
    def enqueue(raw,depth):
        u=normalize(urljoin(start,raw))
        if not u or not same_site(u,family) or urlparse(u).path.lower().endswith(NON_HTML):
            return
        if include_rules and not path_matches_rules(u,include_rules):
            discovery_sources[u].add("excluded_by_include_rules"); return
        if path_matches_rules(u,exclude_rules):
            discovery_sources[u].add("excluded_by_exclude_rules"); return
        if depth is not None:
            old=depths.get(u)
            if old is None or depth < old:
                depths[u]=depth
        if u not in queued and len(queued)<max_pages*50:
            queued.add(u); queue.append(u)
        return u
    enqueue(start,0); discovery_sources[start].add('seed_url')
    if use_sitemap:
        for u in sitemap_urls(session,origin,family,max_pages*20,robots_text,source_diagnostics):
            enqueue(u,None); discovery_sources[u].add('sitemap_or_robots_sitemap')
    for path in (url_files or []):
        for u,src in load_url_inventory(path,start,family,'manual_import:'+str(path)):
            enqueue(u,None); discovery_sources[u].add(src)
    if previous_report:
        for u,src in load_url_inventory(previous_report,start,family,'previous_crawl'):
            enqueue(u,None); discovery_sources[u].add(src)
    for feed in (feed_urls or []):
        for u,src in fetch_xml_or_feed_urls(session,urljoin(origin,feed),family,max_pages*10,source_diagnostics=source_diagnostics,source_kind="feed"):
            enqueue(u,None); discovery_sources[u].add(src)
    source_diagnostics.append({"Source type":"seed URL","Source URL":start,"HTTP status":"","URLs discovered":1,"Source status":"queued","Notes":"Project start URL added to crawl queue"})
    for integration in ("Google Search Console","Google Analytics 4","Backlink provider","CMS inventory","Server logs"):
        source_diagnostics.append({"Source type":integration,"Source URL":"","HTTP status":"","URLs discovered":0,"Source status":"not_connected","Notes":"No authorised integration is configured; no external data was collected or fabricated"})
    if url_files:
        source_diagnostics.append({"Source type":"manual URL import","Source URL":"; ".join(str(x) for x in url_files),"HTTP status":"","URLs discovered":sum(1 for srcs in discovery_sources.values() if any(str(v).startswith('manual_import:') for v in srcs)),"Source status":"processed","Notes":"Only public same-site URLs are accepted"})
    if previous_report:
        source_diagnostics.append({"Source type":"previous crawl","Source URL":str(previous_report),"HTTP status":"","URLs discovered":sum(1 for srcs in discovery_sources.values() if 'previous_crawl' in srcs),"Source status":"processed","Notes":"Imported same-site URLs from prior workbook where available"})
    # A robots.txt sitemap may be cross-host; only same-site URLs are crawled by default.
    for u in queued:
        discovery_rows.append({'URL':u,'Discovery source':' ; '.join(sorted(discovery_sources.get(u,{'unknown'}))),'Source count':len(discovery_sources.get(u,set())),'Discovery status':'queued'})
    last=0
    while queue and len(visited)<max_pages:
        url=queue.popleft()
        if url in visited: continue
        visited.add(url)
        if respect_robots and not robots.can_fetch(ua,url):
            row={h:"" for h in HEADERS}; row.update({"Address":url,"Indexability":"Blocked","Indexability Status":"Disallowed by robots.txt","Crawl Depth":depths.get(url,0),"Crawl Timestamp":utcnow()})
            rows.append(row)
            diagnostics[url]={"Address":url,"status_code":"","content_type":"","html_detected":False,"extraction_status":"blocked_by_robots","body_truncated":False,"raw_bytes_read":0,"declared_bytes":"","rendered_checked":False,"title_tags_raw":0,"meta_description_tags_raw":0,"h1_tags_raw":0,"title_raw":"","meta_description_raw":"","h1_raw":"","title_selected":"","meta_description_selected":"","h1_selected":"","open_graph_description_present":False,"note":"Not fetched because robots.txt disallows this URL"}
            write_progress(progress_path, status="running", crawled=len(visited), max_pages=max_pages, current_url=url, message="Blocked by robots.txt", errors=fetch_errors, queued=len(queue))
            if len(visited)%10==0:
                write_excel(rows,output,diagnostics)
                print(f"PROGRESS {len(visited)} / {max_pages} pages", flush=True)
            continue
        pause=delay-(time.monotonic()-last)
        if pause>0: time.sleep(pause)
        started=time.monotonic(); last=started
        chain=[url]; current=url; response=None; first_status=None; redirect_codes=[]
        try:
            # Request one URL at a time without auto-following redirects. This keeps each
            # report row tied to the HTTP response actually returned for that URL.
            response=session.get(url,timeout=TIMEOUT,allow_redirects=False,stream=True)
            first_status=response.status_code
            if response.is_redirect or response.is_permanent_redirect:
                redirect_codes.append(response.status_code)
                loc=response.headers.get("Location")
                if loc:
                    nxt=normalize(urljoin(url,loc))
                    if nxt:
                        current=nxt
                        chain.append(nxt)
            body=response.raw.read(MAX_BYTES,decode_content=True)
            rendered_body=None; render_error=""
            challenge=is_challenge_response(first_status,body,response.headers)
            if render_js and 200 <= first_status < 300 and not challenge and is_html_response(response.headers.get("Content-Type",""),body):
                rendered_body,render_error=render_page_html(url)
            row,links,diag=extract_page(url,response,url,chain,body,depths.get(url,""),rows,incoming_sources,reported_status=first_status,rendered_body=rendered_body)
            diag["render_error"]=render_error
            if render_error:
                diag["note"] += "; " + render_error
            diag["redirect_chain"]=" -> ".join(chain)
            diag["redirect_status_codes"]=" -> ".join(map(str,redirect_codes))
            diag["response_headers_json"]=json.dumps(dict(response.headers),ensure_ascii=False)
            diagnostics[url]=diag
            try:
                selected_soup=BeautifulSoup(rendered_body or body, "lxml") if 200 <= first_status < 300 and not challenge and is_html_response(response.headers.get("Content-Type",""),rendered_body or body) else None
                extra_rows[url]=extract_extra_page_data(url,url,response,body,selected_soup,links,rendered_body)
                if challenge:
                    extra_rows[url]["Extraction status"]="blocked_challenge"
                    extra_rows[url]["Extraction notes"]="Access challenge detected; SEO content was not assessed. Ask the site owner to allow the authorised crawler."
                extra_rows[url]["Final URL"]="" if redirect_codes else url
                extra_rows[url]["Redirect destination observed"]=current if current != url else ""
                extra_rows[url]["Redirect chain observed"]=" -> ".join(chain)
            except Exception as data_exc:
                extra_rows[url]={"Address":url,"Extraction status":"failed","Extraction error":str(data_exc)}
            response.close()
            row["Response Time"]=round(time.monotonic()-started,3)
            row["Redirect URL"]=current if current != url else ""
            row["Redirect Type"]=" -> ".join(str(code) for code in redirect_codes)
            # A redirect response body is not the destination page. Crawl its same-site
            # destination as a separate URL and do not attribute destination links to source.
            if redirect_codes:
                links=[]
                if same_site(current,family):
                    discovery_sources[current].add("redirect_destination")
                    parent_depth=depths.get(url)
                    enqueue(current,(parent_depth+1) if isinstance(parent_depth,int) else None)
            discovered_links[url]=links
            source_diagnostics.append({"Source type":"rendered internal links" if rendered_body else "raw internal links","Source URL":url,"HTTP status":first_status,"URLs discovered":len(links),"Source status":"checked" if 200 <= first_status < 300 and not challenge else ("blocked" if challenge or first_status in (401,403,429) else "not_applicable"),"Notes":("Links extracted from rendered DOM" if rendered_body else "Links extracted from raw HTML") if 200 <= first_status < 300 and not challenge else ("Access challenge; links not extracted" if challenge else f"HTTP {first_status}; links not extracted")})
            for target in links:
                if same_site(target,family):
                    discovery_sources[target].add("rendered_internal_link" if rendered_body else "raw_internal_link")
                    incoming_total[target] += 1
                    incoming_sources[target].add(url)
                    parent_depth=depths.get(url)
                    enqueue(target,(parent_depth+1) if isinstance(parent_depth,int) else None)
            rows.append(row)
        except requests.RequestException as e:
            fetch_errors += 1
            row={h:"" for h in HEADERS}; row.update({"Address":url,"Status":"Fetch Error","Indexability":"Unknown","Indexability Status":str(e),"Crawl Depth":depths.get(url,0),"Crawl Timestamp":utcnow()}); rows.append(row)
            diagnostics[url]={"Address":url,"status_code":"","content_type":"","html_detected":False,"extraction_status":"fetch_failed","body_truncated":False,"raw_bytes_read":0,"declared_bytes":"","rendered_checked":False,"title_tags_raw":0,"meta_description_tags_raw":0,"h1_tags_raw":0,"title_raw":"","meta_description_raw":"","h1_raw":"","title_selected":"","meta_description_selected":"","h1_selected":"","open_graph_description_present":False,"note":f"Fetch failed: {e}"}
        write_progress(progress_path, status="running", crawled=len(visited), max_pages=max_pages, current_url=url, message="URL processed", errors=fetch_errors, queued=len(queue))
        if len(visited)%10==0:
            write_excel(rows,output,diagnostics)
            print(f"PROGRESS {len(visited)} / {max_pages} pages", flush=True)
    # Recalculate shortest crawl depth from the start URL using the actual internal-link graph.
    # Sitemap-discovered pages that cannot be reached by links keep depth blank.
    page_urls={row.get("Address","") for row in rows if row.get("Address")}
    redirect_targets={row.get("Address",""):row.get("Redirect URL","") for row in rows if row.get("Address") and row.get("Redirect URL")}
    shortest={start:0}; depth_queue=deque([start])
    while depth_queue:
        source=depth_queue.popleft()
        neighbors=list(discovered_links.get(source,[]))
        redirect_target=redirect_targets.get(source,"")
        if redirect_target:
            neighbors.append(redirect_target)
        for target in neighbors:
            if same_site(target,family) and target in page_urls and target not in shortest:
                shortest[target]=shortest[source]+1
                depth_queue.append(target)

    # Calculate a normalized internal-link PageRank-style score for crawled URLs.
    # This is an explanatory internal score, not a claim to reproduce a vendor's proprietary Link Score.
    graph={u:set(t for t in discovered_links.get(u,[]) if t in page_urls and same_site(t,family)) for u in page_urls}
    n=max(len(page_urls),1)
    rank={u:1.0/n for u in page_urls}
    damping=0.85
    for _ in range(20):
        dangling=sum(rank[u] for u in page_urls if not graph.get(u))
        next_rank={u:(1-damping)/n + damping*dangling/n for u in page_urls}
        for source, targets in graph.items():
            if targets:
                share=damping*rank[source]/len(targets)
                for target in targets:
                    next_rank[target]+=share
        rank=next_rank
    max_rank=max(rank.values(),default=1.0)

    # Finalize link counts and exact duplicate-content flags.
    hashes=defaultdict(list)
    for row in rows:
        if row.get("Hash"): hashes[row["Hash"]].append(row)
    total_inlinks=sum(incoming_total.values())
    for row in rows:
        address=row.get("Address","")
        row["Inlinks"]=incoming_total.get(address,0)
        row["Unique Inlinks"]=len(incoming_sources.get(address,set()))
        row["Crawl Depth"]=shortest.get(address,"")
        row["% of Total"]=round(100*row["Inlinks"]/total_inlinks,3) if total_inlinks else 0
        row["Link Score"]=round(100*rank.get(address,0)/max_rank,2) if address in rank and max_rank else ""
        # Retained for export compatibility; only exact body-hash matches are confirmed,
        # not fuzzy/near-duplicate similarity.
        near=hashes.get(row.get("Hash",""),[]) if row.get("Hash") else []
        row["No. Near Duplicates"]=max(0,len(near)-1)
        row["Closest Similarity Match"]=next((x["Address"] for x in near if x["Address"]!=address),"")
        row["CO2 (mg)"]=""; row["Carbon Rating"]=""; row["Unique JS Inlinks"]=""; row["Unique JS Outlinks"]=""; row["Unique External JS Outlinks"]=""
    # Finalize URL discovery records, including known URLs that were queued but not reached before max_pages.
    discovery_by_url={r.get('URL'):r for r in discovery_rows}
    for u,sources in discovery_sources.items():
        if u not in discovery_by_url: discovery_rows.append({'URL':u,'Discovery source':' ; '.join(sorted(sources)),'Source count':len(sources),'Discovery status':'crawled' if u in visited else 'queued_not_crawled'})
        else:
            discovery_by_url[u]['Discovery source']=' ; '.join(sorted(sources)); discovery_by_url[u]['Source count']=len(sources); discovery_by_url[u]['Discovery status']='crawled' if u in visited else 'queued_not_crawled'
    crawled_set={r.get('Address') for r in rows}
    incoming_set=set(incoming_sources)
    for rec in discovery_rows:
        rec['Orphan candidate']='yes' if rec.get('URL') in crawled_set and rec.get('URL')!=start and rec.get('URL') not in incoming_set else ('not_assessed' if rec.get('URL') not in crawled_set else 'no')
    successful_pages=sum(1 for r in rows if 200 <= int(r.get("Status Code") or 0) < 300 and diagnostics.get(r.get("Address"),{}).get("extraction_status")=="checked")
    blocked_pages=sum(1 for r in rows if diagnostics.get(r.get("Address"),{}).get("extraction_status")=="blocked_challenge" or int(r.get("Status Code") or 0) in (401,403,429,503) or str(r.get("Indexability Status","")).startswith("Disallowed by robots"))
    crawl_summary=[{"Metric":"Start URL","Value":start},{"Metric":"URLs discovered/queued","Value":len(queued)},{"Metric":"URLs attempted","Value":len(visited)},{"Metric":"Successful pages with checked extraction","Value":successful_pages},{"Metric":"Blocked/challenge URLs","Value":blocked_pages},{"Metric":"Fetch errors","Value":fetch_errors},{"Metric":"Max URL limit","Value":max_pages},{"Metric":"robots.txt status","Value":robots_state},{"Metric":"Respect robots.txt","Value":bool(respect_robots)},{"Metric":"JavaScript rendering requested","Value":bool(render_js)},{"Metric":"User-Agent","Value":ua},{"Metric":"Crawl quality","Value":"incomplete_blocked" if blocked_pages and not successful_pages else ("usable_partial" if blocked_pages else ("usable" if successful_pages else "no_successful_extractions"))}]
    write_excel(rows,output,diagnostics,extra_rows,discovery_rows,source_diagnostics,crawl_summary)
    write_progress(progress_path, status="finished", crawled=len(visited), max_pages=max_pages, current_url="", message=f"Finished: {len(rows)} URLs; {successful_pages} successfully extracted; {blocked_pages} blocked/challenge URLs", errors=fetch_errors, queued=0)
    print(f"FINISHED {len(rows)} pages. Excel report: {output}", flush=True)
    return len(rows)

def main():
    p=argparse.ArgumentParser()
    p.add_argument("url"); p.add_argument("--max-pages",type=int,default=500)
    p.add_argument("--delay",type=float,default=1.0); p.add_argument("--output",default="seo_audit_report.xlsx")
    p.add_argument("--no-sitemap",action="store_true")
    p.add_argument("--progress",default="")
    p.add_argument("--render-js",action="store_true",help="Render pages in Chromium (requires Playwright and Chromium installed)")
    p.add_argument("--url-file",action="append",default=[],help="Import additional URLs from .txt, .csv, or .xlsx; repeatable")
    p.add_argument("--previous-report",default="",help="Import URLs from a previous Rankpilot/Excel crawl report")
    p.add_argument("--feed",action="append",default=[],help="Additional same-site RSS/Atom feed URL; repeatable")
    p.add_argument("--ignore-robots",action="store_true",help="Ignore robots.txt only when explicitly authorised by the site owner")
    p.add_argument("--user-agent",default=UA,help="Crawler user-agent string")
    p.add_argument("--include",action="append",default=[],help="Include URL path/full-URL glob; repeatable")
    p.add_argument("--exclude",action="append",default=[],help="Exclude URL path/full-URL glob; repeatable")
    a=p.parse_args()
    if not 1<=a.max_pages<=100000: p.error("--max-pages must be between 1 and 100000")
    if a.delay<0.5: p.error("--delay must be at least 0.5 seconds")
    try:
        run_crawl(a.url,a.max_pages,a.delay,a.output,not a.no_sitemap,a.progress or None,a.render_js,a.url_file,a.previous_report or None,a.feed,not a.ignore_robots,a.user_agent,a.include,a.exclude)
    except Exception as exc:
        write_progress(a.progress or None, status="failed", crawled=0, max_pages=a.max_pages, message=str(exc))
        raise
if __name__=="__main__": main()

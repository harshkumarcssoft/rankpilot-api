from datetime import timedelta
from io import BytesIO
from pathlib import Path
import tempfile

from openpyxl import load_workbook

import seo_crawler


class FakeRaw:
    def __init__(self, body):
        self.body = body
        self.version = 11

    def read(self, size=-1, decode_content=False):
        data = self.body if size is None or size < 0 else self.body[:size]
        self.body = b"" if size is None or size < 0 else self.body[size:]
        return data


class FakeResponse:
    def __init__(self, url, body, status=200, content_type="text/html; charset=utf-8", headers=None):
        self.url = url
        self.status_code = status
        self.headers = {"Content-Type": content_type, "Content-Length": str(len(body)), **(headers or {})}
        self._content = body
        self.raw = FakeRaw(body)
        self.elapsed = timedelta(milliseconds=25)
        self.is_redirect = 300 <= status < 400
        self.is_permanent_redirect = status in (301, 308)

    @property
    def text(self):
        return self._content.decode("utf-8", errors="replace")

    def close(self):
        pass


class FakeSession:
    def __init__(self):
        self.headers = {}

    def get(self, url, **kwargs):
        if url.endswith("/robots.txt"):
            return FakeResponse(url, b"User-agent: *\nAllow: /\nSitemap: https://example.test/sitemap.xml\n", content_type="text/plain")
        if url.endswith("/sitemap.xml"):
            xml = b'<?xml version="1.0"?><urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9"><url><loc>https://example.test/</loc></url><url><loc>https://example.test/about</loc></url></urlset>'
            return FakeResponse(url, xml, content_type="application/xml")
        if url.endswith("/sitemap_index.xml"):
            return FakeResponse(url, b"not found", status=404, content_type="text/plain")
        pages = {
            "https://example.test/": b'<html lang="en"><head><title>Home</title><meta name="description" content="Home summary"></head><body><h1>Home</h1><a href="/about">About</a><a href="/contact">Contact</a></body></html>',
            "https://example.test/about": b'<html lang="en"><head><title>About</title><meta name="description" content="About summary"></head><body><h1>About</h1><a href="/contact">Contact</a></body></html>',
            "https://example.test/contact": b'<html lang="en"><head><title>Contact</title><meta name="description" content="Contact summary"></head><body><h1>Contact</h1></body></html>',
        }
        if url in pages:
            return FakeResponse(url, pages[url])
        return FakeResponse(url, b"not found", status=404, content_type="text/plain")


def test_crawler_discovers_sitemap_and_recursive_internal_links(monkeypatch):
    monkeypatch.setattr(seo_crawler, "public_host", lambda host: True)
    monkeypatch.setattr(seo_crawler.requests, "Session", FakeSession)
    monkeypatch.setattr(seo_crawler.time, "sleep", lambda seconds: None)
    with tempfile.TemporaryDirectory() as td:
        output = Path(td) / "crawl.xlsx"
        progress = Path(td) / "progress.json"
        count = seo_crawler.run_crawl(
            "https://example.test/", 20, 0.5, str(output), True, str(progress), False,
            [], None, [], True, "TestCrawler/1.0", [], []
        )
        assert count == 3
        wb = load_workbook(output, read_only=True, data_only=True)
        ws = wb["1 - All"]
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        address_col = headers.index("Address")
        assert {row[address_col] for row in rows} == {
            "https://example.test/", "https://example.test/about", "https://example.test/contact"
        }
        assert {"URL Discovery", "Discovery Diagnostics", "Crawl Summary", "Page Data"}.issubset(set(wb.sheetnames))
        diag = list(wb["Discovery Diagnostics"].iter_rows(values_only=True))
        assert any(row and row[0] == "raw internal links" for row in diag[1:])
        summary = list(wb["Crawl Summary"].iter_rows(values_only=True))
        summary_dict = {row[0]: row[1] for row in summary[1:]}
        assert summary_dict["Successful pages with checked extraction"] == 3
        assert summary_dict["Crawl quality"] == "usable"
        wb.close()

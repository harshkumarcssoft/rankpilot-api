import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from requests import Response
from seo_crawler import extract_page, HEADERS, write_excel
from openpyxl import load_workbook
import tempfile


def fake_response(status=200, content_type='text/html; charset=utf-8'):
    r=Response(); r.status_code=status; r.headers['Content-Type']=content_type
    r.headers['Content-Length']='0'; r._content=b''; r.url='https://example.test/'; r.elapsed=__import__('datetime').timedelta(milliseconds=25)
    return r


def test_first_empty_description_does_not_hide_later_populated_description():
    html=b'<!doctype html><html><head><title>Test page</title><meta name="description" content="   "><META NAME="Description" CONTENT="A useful page summary"></head><body><h1>Heading</h1><a href="/next">Next</a></body></html>'
    row, links, diag=extract_page('https://example.test/', fake_response(), 'https://example.test/', ['https://example.test/'], html, 0, [], {}, reported_status=200)
    assert row['Meta Description 1']=='A useful page summary'
    assert diag['meta_description_tags_raw']==2
    assert diag['extraction_status']=='checked'


def test_html_is_detected_when_content_type_is_missing():
    html=b'<html><head><title>Present title</title><meta name="description" content="Present description"></head><body></body></html>'
    row, links, diag=extract_page('https://example.test/', fake_response(content_type=''), 'https://example.test/', ['https://example.test/'], html, 0, [], {}, reported_status=200)
    assert row['Title 1']=='Present title'
    assert row['Meta Description 1']=='Present description'
    assert diag['html_detected'] is True


def test_error_response_is_not_eligible_for_missing_tag_diagnosis():
    html=b'<html><head><title>Error page</title></head><body>Not found</body></html>'
    row, links, diag=extract_page('https://example.test/missing', fake_response(status=404), 'https://example.test/missing', ['https://example.test/missing'], html, 0, [], {}, reported_status=404)
    assert diag['extraction_status']=='not_applicable'


def test_export_includes_all_dynamic_fields_without_a_fixed_column_limit():
    html=b'<html><head><title>Present</title><meta name="description" content="Present"></head><body><h1>Heading</h1></body></html>'
    row, links, diag=extract_page('https://example.test/', fake_response(), 'https://example.test/', ['https://example.test/'], html, 0, [], {}, reported_status=200)
    row['Future crawler field'] = 'included'
    extra={'https://example.test/': {'All headings JSON': '["Heading"]', 'Future page field': 'page-value'}}
    with tempfile.TemporaryDirectory() as td:
        out=Path(td)/'audit.xlsx'; write_excel([row],out,{'https://example.test/':diag},extra)
        wb=load_workbook(out,read_only=True,data_only=True)
        assert '1 - All' in wb.sheetnames and 'Audit Diagnostics' in wb.sheetnames and 'Page Data' in wb.sheetnames
        ws=wb['1 - All']; values=list(ws.iter_rows(values_only=True)); headers=list(values[0]); data=dict(zip(headers,values[1]))
        assert len(headers) > len(HEADERS)
        assert data['Future crawler field']=='included'
        assert data['Page Data: All headings JSON']=='["Heading"]'
        assert data['Page Data: Future page field']=='page-value'
        dws=wb['Audit Diagnostics']; dheaders=[c.value for c in next(dws.iter_rows(min_row=1,max_row=1))]
        assert 'extraction_status' in dheaders and 'meta_description_tags_raw' in dheaders
        wb.close()


def test_extended_page_data_collects_meta_headings_images_and_security_headers():
    from seo_crawler import extract_extra_page_data
    html=b'''<!doctype html><html lang="en"><head><title>Page title</title><meta name="description" content="Summary"><meta property="og:title" content="Social title"><meta name="viewport" content="width=device-width, initial-scale=1"><link rel="canonical" href="https://example.test/canonical"><script type="application/ld+json">{"@type":"LocalBusiness"}</script></head><body><h1>Main</h1><p>Hello world.</p><img src="/photo.jpg" alt="Photo" width="100"><a href="/next">Next page</a><form><input id="email"><label for="email">Email</label></form></body></html>'''
    r=fake_response(); r.headers['Strict-Transport-Security']='max-age=31536000'
    from bs4 import BeautifulSoup
    soup=BeautifulSoup(html,'lxml')
    data=extract_extra_page_data('https://example.test/','https://example.test/',r,html,soup,[])
    assert data['HTML lang']=='en'
    assert data['H1 count']==1
    assert 'Summary' in data['Meta tags JSON']
    assert 'photo.jpg' in data['Image data JSON']
    assert 'LocalBusiness' in data['Structured data JSON']
    assert 'max-age=31536000' in data['Response security headers JSON']


def test_url_inventory_imports_same_site_text_urls_and_rejects_other_hosts(monkeypatch):
    import seo_crawler
    monkeypatch.setattr(seo_crawler, "public_host", lambda host: True)
    from seo_crawler import load_url_inventory
    import tempfile
    with tempfile.TemporaryDirectory() as td:
        f=Path(td)/'urls.txt'
        f.write_text('https://example.com/one\nhttps://other.com/two\n/relative-path\n',encoding='utf-8')
        found=load_url_inventory(f,'https://example.com/','example.com','manual')
        urls={u for u,_ in found}
        assert 'https://example.com/one' in urls
        assert 'https://example.com/relative-path' in urls
        assert 'https://other.com/two' not in urls


def test_cloudflare_challenge_is_not_treated_as_real_page_content():
    from seo_crawler import is_challenge_response
    html = b'<!doctype html><html><head><title>Just a moment...</title></head><body><div id="challenge-platform">Checking your browser</div><a href="/real-page">Continue</a></body></html>'
    response = fake_response(status=403)
    response.headers['Server'] = 'cloudflare'
    row, links, diag = extract_page(
        'https://example.test/', response, 'https://example.test/', ['https://example.test/'],
        html, 0, [], {}, reported_status=403,
    )
    assert is_challenge_response(403, html, response.headers) is True
    assert links == []
    assert row['Title 1'] == ''
    assert row['Indexability'] == 'Blocked'
    assert diag['extraction_status'] == 'blocked_challenge'
    assert 'no seo metadata or links were extracted' in diag['note'].lower()


def test_url_include_exclude_rules_are_applied_to_paths():
    from seo_crawler import path_matches_rules
    assert path_matches_rules('https://example.test/blog/post-a', ['/blog/*'])
    assert path_matches_rules('https://example.test/checkout/cart', ['/checkout/*'])
    assert not path_matches_rules('https://example.test/about', ['/blog/*'])


def test_workbook_always_has_discovery_diagnostics_and_crawl_summary_sheets():
    html=b'<html><head><title>Present</title><meta name="description" content="Present"></head><body><h1>Heading</h1></body></html>'
    # A real requests.Response is required by extract_page for response metadata.
    response=fake_response()
    row, links, diag=extract_page('https://example.test/',response,'https://example.test/',[],html,0,[],{},reported_status=200)
    with tempfile.TemporaryDirectory() as td:
        out=Path(td)/'audit.xlsx'
        write_excel([row],out,{'https://example.test/':diag},{'https://example.test/':{'Address':'https://example.test/','Visible text':'Heading'}},[{'URL':'https://example.test/','Discovery source':'seed_url'}],[],[])
        wb=load_workbook(out,read_only=True,data_only=True)
        assert 'Discovery Diagnostics' in wb.sheetnames
        assert 'Crawl Summary' in wb.sheetnames
        assert next(wb['Discovery Diagnostics'].iter_rows(values_only=True))[0] == 'Source type'
        assert next(wb['Crawl Summary'].iter_rows(values_only=True))[0] == 'Metric'
        wb.close()
